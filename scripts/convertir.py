"""
Convierte la clasificación (clasificacion.xlsx) y los experimentos (experimentos/*.html)
en los archivos que descarga la app:

  public/version.json   -> {"version": "..."}  (la app lo consulta para saber si hay cambios)
  public/paquete.json   -> catálogo completo con el contenido de cada experimento
  public/img/...        -> imágenes usadas en los experimentos

Se ejecuta automáticamente en GitHub cada vez que subes cambios.
Para probarlo en tu computadora:  pip install openpyxl beautifulsoup4
                                  python scripts/convertir.py
"""
import hashlib
import json
import re
import shutil
import sys
import unicodedata

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from datetime import datetime, timezone
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
EXCEL = RAIZ / "clasificacion.xlsx"
CARPETA_HTML = RAIZ / "experimentos"
CARPETA_IMG = RAIZ / "imagenes"
SALIDA = RAIZ / "public"

FILA_ENCABEZADOS = 2
PRIMERA_FILA_DATOS = 4  # la fila 3 es el ejemplo de la plantilla

avisos: list[str] = []


def clave(texto: str) -> str:
    """'Almacenamiento de energía' -> 'almacenamiento-de-energia' (igual que en la app)."""
    t = unicodedata.normalize("NFD", str(texto))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^a-z0-9]+", "-", t.lower())
    return t.strip("-")


def texto_celda(v) -> str:
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------- Excel
def leer_excel() -> list[dict]:
    if not EXCEL.exists():
        sys.exit(f"ERROR: no se encontró {EXCEL.name} en la raíz del repositorio.")
    ws = load_workbook(EXCEL, data_only=True)["Experimentos"]

    # Grupo de cada columna (fila 1, celdas combinadas) y su encabezado (fila 2)
    grupo_actual, columnas = "", []
    for c in range(1, ws.max_column + 1):
        g = texto_celda(ws.cell(1, c).value)
        if g:
            grupo_actual = g.upper()
        columnas.append((grupo_actual, texto_celda(ws.cell(FILA_ENCABEZADOS, c).value)))

    def es_grupo(g: str, palabra: str) -> bool:
        return palabra in g

    # Columna opcional "Referencia" (puede estar en cualquier lugar de la hoja)
    col_ref = next((i for i, (_, h) in enumerate(columnas) if clave(h) in ("referencia", "referencias")), None)

    filas = []
    for r in range(PRIMERA_FILA_DATOS, ws.max_row + 1):
        valores = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        datos = {h: v for (g, h), v in zip(columnas, valores) if g == "DATOS"}
        archivo = texto_celda(datos.get("Archivo HTML"))
        if not archivo:
            continue
        marcados = lambda palabra: [
            clave(h) for (g, h), v in zip(columnas, valores)
            if es_grupo(g, palabra) and texto_celda(v).lower() == "x"
        ]
        fila = {
            "fila": r,
            "archivo": archivo,
            "titulo": texto_celda(datos.get("Título")),
            "resumen": texto_celda(datos.get("Resumen")),
            "nivel": texto_celda(datos.get("Nivel")),
            "duracionMin": datos.get("Duración (min)") or None,
            "kits": [k.strip() for k in texto_celda(datos.get("Kits")).split(",") if k.strip()],
            "nobel": datos.get("Nobel (año)") or None,
            "topicos": marcados("TÓPICOS"),
            "campos": marcados("CAMPOS"),
            "seguridad": marcados("SEGURIDAD"),
            "referencias": leer_referencias(valores[col_ref]) if col_ref is not None else [],
        }
        for campo_num in ("duracionMin", "nobel"):
            try:
                fila[campo_num] = int(fila[campo_num]) if fila[campo_num] else None
            except (TypeError, ValueError):
                avisos.append(f"Fila {r}: '{campo_num}' no es un número; se ignora.")
                fila[campo_num] = None
        if not fila["topicos"] and not fila["campos"]:
            avisos.append(f"Fila {r} ({archivo}): no tiene ningún tópico ni campo marcado con 'x'.")
        filas.append(fila)
    return filas


URL = re.compile(r"https?://[^\s,;<>\"')]+[^\s,;<>\"').]")


def leer_referencias(valor) -> list[dict]:
    """Enlaces y/o citas separados por salto de línea o punto y coma.
    Cada enlace se vuelve una referencia; un texto sin enlace se guarda como cita."""
    refs = []
    for parte in re.split(r"[\n;]+", texto_celda(valor)):
        parte = parte.strip()
        if not parte:
            continue
        urls = URL.findall(parte)
        if not urls:
            refs.append({"texto": parte, "url": None})
            continue
        cita = limpiar(URL.sub("", parte)).strip(" ,.-–:")
        for u in urls:
            if cita and len(urls) == 1:
                texto = cita
            else:
                texto = re.sub(r"^https?://(www\.)?", "", u).rstrip("/")
                texto = texto if len(texto) <= 60 else texto[:57] + "…"
            refs.append({"texto": texto, "url": u})
    return refs


# ---------------------------------------------------------------- HTML
SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
SUP = str.maketrans("0123456789+-=()", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾")


def texto_inline(nodo) -> str:
    """Texto de un nodo conservando subíndices y superíndices (H₂O, cm³)."""
    if isinstance(nodo, NavigableString):
        return str(nodo)
    if not isinstance(nodo, Tag):
        return ""
    if nodo.name == "sub":
        return nodo.get_text().translate(SUB)
    if nodo.name == "sup":
        return nodo.get_text().translate(SUP)
    if nodo.name == "br":
        return "\n"
    return "".join(texto_inline(h) for h in nodo.children)


def limpiar(t: str) -> str:
    return re.sub(r"\s+", " ", t).strip()


def lineas_de_parrafo(p: Tag) -> list[tuple[str, bool]]:
    """Divide un <p> por los <br>. Devuelve (texto, es_aviso) — aviso = línea toda en negrita."""
    segmentos, actual = [], []
    for h in p.children:
        if isinstance(h, Tag) and h.name == "br":
            segmentos.append(actual)
            actual = []
        else:
            actual.append(h)
    segmentos.append(actual)

    resultado = []
    for seg in segmentos:
        texto = limpiar("".join(texto_inline(n) for n in seg))
        if not texto:
            continue
        etiquetas = [n for n in seg if isinstance(n, Tag)]
        sueltos = "".join(str(n) for n in seg if isinstance(n, NavigableString)).strip()
        aviso = bool(etiquetas) and not sueltos and all(n.name in ("strong", "b", "em", "i") for n in etiquetas) \
            and any(n.name in ("strong", "b") for n in etiquetas)
        resultado.append((texto, aviso))
    return resultado


def tipo_seccion(titulo: str) -> str:
    t = clave(titulo)
    if "material" in t or "equipo" in t:
        return "materiales"
    if "procedimiento" in t or "metodo" in t or "pasos" in t:
        return "pasos"
    if "objetivo" in t:
        return "objetivos"
    if "discusion" in t or "conclusion" in t or "pregunta" in t or "cuestionario" in t:
        return "preguntas"
    return "texto"


def titulo_corto(titulo: str) -> str:
    # "IV. Procedimiento" -> "Procedimiento"
    return re.sub(r"^\s*([IVXLC]+|\d+)[\.\)\-]\s*", "", titulo).strip()


def copiar_imagen(src: str, html: Path) -> str | None:
    if not src or src.startswith(("http://", "https://", "data:")):
        return src or None
    candidatos = [html.parent / src, CARPETA_IMG / src, CARPETA_IMG / Path(src).name]
    origen = next((c for c in candidatos if c.is_file()), None)
    if origen is None:
        avisos.append(f"{html.name}: no se encontró la imagen '{src}' (búscala junto al HTML o en /imagenes).")
        return None
    destino_nombre = f"{html.stem}__{origen.name}"  # evita choques entre experimentos (picture.png)
    (SALIDA / "img").mkdir(parents=True, exist_ok=True)
    shutil.copy2(origen, SALIDA / "img" / destino_nombre)
    return f"img/{destino_nombre}"


def leer_html(ruta: Path) -> dict:
    sopa = BeautifulSoup(ruta.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for basura in sopa.select(".botones, script, style, button"):
        basura.decompose()

    titulo_el = sopa.select_one(".titulo") or sopa.find("h1")
    titulo = limpiar(texto_inline(titulo_el)) if titulo_el else limpiar(sopa.title.get_text() if sopa.title else ruta.stem)

    secciones: list[dict] = []

    def seccion_actual() -> dict:
        if not secciones:
            secciones.append({"titulo": "Presentación", "tipo": "texto", "bloques": []})
        return secciones[-1]

    cuerpo = sopa.body or sopa
    for el in cuerpo.find_all(["div", "h2", "h3", "p", "li", "img"]):
        clases = el.get("class") or []
        if el.name in ("h2", "h3") or (el.name == "div" and "subtitulo" in clases):
            t = titulo_corto(limpiar(texto_inline(el)))
            secciones.append({"titulo": t, "tipo": tipo_seccion(t), "bloques": []})
        elif el.name in ("p", "li"):
            if el.find_parent("li") and el.name == "p":
                continue
            sec = seccion_actual()
            lineas = lineas_de_parrafo(el) if el.name == "p" else [(limpiar(texto_inline(el)), False)]
            for texto, aviso in lineas:
                if aviso:
                    sec["bloques"].append({"t": "aviso", "texto": texto})
                elif texto.startswith(("•", "-", "–", "·")) or el.name == "li":
                    sec["bloques"].append({"t": "item", "texto": texto.lstrip("•-–· ").strip()})
                elif sec["tipo"] == "preguntas" and texto.rstrip().endswith("?"):
                    sec["bloques"].append({"t": "item", "texto": texto})
                else:
                    sec["bloques"].append({"t": "p", "texto": texto})
        elif el.name == "img":
            src = copiar_imagen(el.get("src", ""), ruta)
            if src:
                seccion_actual()["bloques"].append({"t": "img", "src": src, "alt": el.get("alt", "")})

    secciones = [s for s in secciones if s["bloques"]]
    if not secciones:
        avisos.append(f"{ruta.name}: no se encontró contenido (¿usa la estructura .subtitulo / <p>?).")
    return {"titulo": titulo, "secciones": secciones}


# ---------------------------------------------------------------- Principal
def main():
    if SALIDA.exists():
        shutil.rmtree(SALIDA)
    SALIDA.mkdir(parents=True)

    filas = leer_excel()
    html_disponibles = {p.name.lower(): p for p in CARPETA_HTML.glob("*.htm*")}
    usados = set()
    experimentos = []
    ids = set()

    for f in filas:
        ruta = html_disponibles.get(f["archivo"].lower())
        if ruta is None:
            avisos.append(f"Fila {f['fila']}: no existe experimentos/{f['archivo']} — se omite.")
            continue
        usados.add(ruta.name.lower())
        contenido = leer_html(ruta)
        exp_id = clave(ruta.stem)
        if exp_id in ids:
            avisos.append(f"Fila {f['fila']}: {f['archivo']} está repetido — se omite.")
            continue
        ids.add(exp_id)

        resumen = f["resumen"]
        if not resumen:
            primer_p = next((b["texto"] for s in contenido["secciones"] for b in s["bloques"] if b["t"] == "p"), "")
            resumen = (primer_p[:180] + "…") if len(primer_p) > 180 else primer_p

        experimentos.append({
            "id": exp_id,
            "archivo": ruta.name,
            "titulo": f["titulo"] or contenido["titulo"],
            "resumen": resumen,
            "nivel": f["nivel"] or None,
            "duracionMin": f["duracionMin"],
            "kits": f["kits"],
            "nobel": f["nobel"],
            "topicos": f["topicos"],
            "campos": f["campos"],
            "seguridad": f["seguridad"],
            "referencias": f["referencias"],
            "secciones": contenido["secciones"],
        })

    for nombre in sorted(set(html_disponibles) - usados):
        avisos.append(f"experimentos/{html_disponibles[nombre].name} no está en el Excel — no se publica.")

    experimentos.sort(key=lambda e: e["titulo"].lower())
    print(f"✔ {len(experimentos)} experimentos")

    # Realidad Aumentada (ra.xlsx), si existe
    import convertir_ra
    ra = convertir_ra.generar(SALIDA, avisos)

    cuerpo = json.dumps({"experimentos": experimentos, "ra": ra}, ensure_ascii=False, sort_keys=True)
    version = hashlib.sha256(cuerpo.encode()).hexdigest()[:12]

    paquete = {
        "version": version,
        "generado": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "experimentos": experimentos,
        "ra": ra,
    }
    (SALIDA / "paquete.json").write_text(json.dumps(paquete, ensure_ascii=False), encoding="utf-8")
    (SALIDA / "version.json").write_text(json.dumps({"version": version, "total": len(experimentos)}), encoding="utf-8")
    (SALIDA / "index.html").write_text(
        f"<!doctype html><meta charset='utf-8'><title>Illari – experimentos</title>"
        f"<h1>Illari OpenLab Académico</h1><p>{len(experimentos)} experimentos · "
        f"{len(ra['temas']) if ra else 0} temas de RA · versión {version}</p>"
        f"<p><a href='paquete.json'>paquete.json</a> · <a href='version.json'>version.json</a></p>",
        encoding="utf-8",
    )

    print(f"\n✔ Publicado (versión {version})")
    if avisos:
        print(f"\n⚠ {len(avisos)} aviso(s):")
        for a in avisos:
            print("  -", a)


if __name__ == "__main__":
    main()
