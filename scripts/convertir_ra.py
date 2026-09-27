"""
Realidad Aumentada: convierte ra.xlsx + modelos .glb en lo que usan la app y los visores.

  public/ra/temas.json             catálogo de temas (lo usa el visor web)
  public/ra/modelos/*.glb          modelos comprimidos (Draco) — 10 a 100 veces más livianos
  public/ra/marcadores/*.png       imagen de cada marcador (A, B, C…)
  public/ra/marcadores/guia_marcadores.pdf   hoja para imprimir
  public/ra/visor.html, visor3d.html          visores (se copian desde web/ra/)

Lo llama convertir.py; no hace falta ejecutarlo por separado.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import struct
import subprocess
import time
import urllib.error
import urllib.request
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
EXCEL_RA = RAIZ / "ra.xlsx"
CARPETA_RA = RAIZ / "ra"
WEB_RA = RAIZ / "web" / "ra"
CACHE = RAIZ / ".cache-modelos"

# Si un modelo no está en ra/modelos/, se busca aquí (en este orden)
FUENTES_MODELOS = [
    "https://raw.githubusercontent.com/lvaapq/asdfgh/main/",
]

# Tipo de código de barras AR (debe coincidir con las imágenes de ra/marcadores/codigos/)
TIPO_CODIGO = "3x3_PARITY65"

GLTF_TRANSFORM = os.environ.get("GLTF_TRANSFORM", "gltf-transform")
OPCIONES_COMPRESION = ["--compress", "draco", "--texture-compress", "false", "--instance", "false"]


def _clave(texto: str) -> str:
    from convertir import clave
    return clave(texto)


def _texto(v) -> str:
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------- Marcadores
def leer_marcadores(ws, avisos: list[str]) -> tuple[dict, str]:
    guia = _texto(ws["B1"].value)
    marcadores: dict[str, dict] = {}
    for r in range(4, ws.max_row + 1):
        mid = _texto(ws.cell(r, 1).value).upper()
        if not mid:
            continue
        tipo = _clave(_texto(ws.cell(r, 2).value) or "codigo")
        valor = _texto(ws.cell(r, 3).value)
        imagen = _texto(ws.cell(r, 4).value)
        if tipo.startswith("cod"):
            try:
                codigo = int(float(valor))
                assert 0 <= codigo <= 31
            except Exception:
                avisos.append(f"Marcadores, fila {r}: el código de '{mid}' debe ser un número entre 0 y 31.")
                continue
            marcadores[mid] = {"id": mid, "tipo": "codigo", "codigo": codigo}
        else:
            if not valor.lower().endswith(".patt"):
                avisos.append(f"Marcadores, fila {r}: el patrón de '{mid}' debe ser un archivo .patt.")
                continue
            marcadores[mid] = {"id": mid, "tipo": "patron", "patronFuente": valor, "imagenFuente": imagen}
    return marcadores, guia


# ---------------------------------------------------------------- Modelos
def obtener_archivo(nombre: str, subcarpeta: str) -> tuple[bytes | None, str]:
    """Busca un archivo en ra/<subcarpeta>/, luego en las fuentes en línea."""
    if nombre.startswith(("http://", "https://")):
        urls = [nombre]
    else:
        local = CARPETA_RA / subcarpeta / nombre
        if local.is_file():
            return local.read_bytes(), str(local.relative_to(RAIZ))
        urls = [f + urllib.request.quote(nombre) for f in FUENTES_MODELOS]
    for url in urls:
        for intento in range(3):  # reintenta si la conexión falla
            try:
                with urllib.request.urlopen(url, timeout=120) as r:
                    return r.read(), url
            except urllib.error.HTTPError as e:
                if e.code == 404:
                    break  # no existe en esta fuente: probar la siguiente
            except Exception:
                pass
            time.sleep(2 * (intento + 1))
    return None, ""


def contar_animaciones(glb: bytes) -> int:
    try:
        largo_json = struct.unpack_from("<I", glb, 12)[0]
        datos = json.loads(glb[20:20 + largo_json])
        return len(datos.get("animations", []))
    except Exception:
        return 0


def comprimir(original: bytes, avisos: list[str], nombre: str) -> bytes:
    """Comprime con gltf-transform (Draco). Usa caché para no repetir trabajo."""
    CACHE.mkdir(exist_ok=True)
    firma = hashlib.sha256(original + " ".join(OPCIONES_COMPRESION).encode()).hexdigest()[:20]
    en_cache = CACHE / f"{firma}.glb"
    if en_cache.exists():
        return en_cache.read_bytes()
    entrada = CACHE / f"{firma}-in.glb"
    entrada.write_bytes(original)
    try:
        subprocess.run(
            [GLTF_TRANSFORM, "optimize", str(entrada), str(en_cache), *OPCIONES_COMPRESION],
            check=True, capture_output=True, timeout=600,
        )
        resultado = en_cache.read_bytes()
        if len(resultado) >= len(original):
            en_cache.write_bytes(original)
            return original
        return resultado
    except Exception as e:
        avisos.append(f"No se pudo comprimir {nombre} ({type(e).__name__}); se publica sin comprimir.")
        return original
    finally:
        entrada.unlink(missing_ok=True)


class Publicador:
    def __init__(self, destino: Path, avisos: list[str]):
        self.destino = destino
        self.avisos = avisos
        self.hechos: dict[str, dict] = {}
        (destino / "modelos").mkdir(parents=True, exist_ok=True)

    def modelo(self, archivo: str) -> dict | None:
        if archivo in self.hechos:
            return self.hechos[archivo]
        datos, origen = obtener_archivo(archivo, "modelos")
        if datos is None:
            self.avisos.append(f"No se encontró el modelo '{archivo}' (ni en ra/modelos/ ni en el repositorio de modelos).")
            self.hechos[archivo] = None
            return None
        comprimido = comprimir(datos, self.avisos, archivo)
        nombre = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(archivo.split("?")[0]).stem) + ".glb"
        (self.destino / "modelos" / nombre).write_bytes(comprimido)
        info = {
            "archivo": f"ra/modelos/{nombre}",
            "kb": round(len(comprimido) / 1024),
            "kbOriginal": round(len(datos) / 1024),
            "animado": contar_animaciones(comprimido) > 0,
        }
        print(f"   modelo {archivo}: {info['kbOriginal']:,} KB → {info['kb']:,} KB")
        self.hechos[archivo] = info
        return info


# ---------------------------------------------------------------- Imágenes y PDF de marcadores
def imagen_marcador_codigo(codigo: int, letra: str, destino: Path):
    from PIL import Image, ImageDraw, ImageFont
    fuente = Image.open(CARPETA_RA / "marcadores" / "codigos" / f"{codigo}.png").convert("L")
    lado, margen, pie = 480, 60, 110
    marca = fuente.resize((lado, lado), Image.NEAREST)
    lienzo = Image.new("L", (lado + 2 * margen, lado + 2 * margen + pie), 255)
    lienzo.paste(marca, (margen, margen))
    d = ImageDraw.Draw(lienzo)
    try:
        f = ImageFont.load_default(size=80)
    except TypeError:
        f = ImageFont.load_default()
    d.text((lienzo.width / 2, lado + 2 * margen + pie / 2 - 20), letra, fill=0, font=f, anchor="mm")
    lienzo.save(destino, optimize=True)


def pdf_marcadores(marcadores: list[dict], carpeta: Path) -> Path | None:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.pdfgen import canvas
    except ImportError:
        return None
    ruta = carpeta / "guia_marcadores.pdf"
    c = canvas.Canvas(str(ruta), pagesize=A4)
    ancho, alto = A4
    lado = 7 * cm
    por_pagina = 4
    for i, m in enumerate(marcadores):
        if i % por_pagina == 0:
            if i:
                c.showPage()
            c.setFont("Helvetica-Bold", 14)
            c.drawCentredString(ancho / 2, alto - 1.6 * cm, "Illari OpenLab Académico · Marcadores de Realidad Aumentada")
            c.setFont("Helvetica", 9)
            c.drawCentredString(ancho / 2, alto - 2.2 * cm,
                                "Imprime al 100 % (sin 'ajustar a la página'). Deja el borde blanco al recortar. Papel mate si es posible.")
        k = i % por_pagina
        col, fila = k % 2, k // 2
        cx = ancho / 4 + col * ancho / 2
        cy = alto - 3.2 * cm - (fila + 0.5) * (alto - 4.5 * cm) / 2
        img = carpeta / f"{m['id']}.png"
        if m["tipo"] == "codigo":
            from PIL import Image
            png = Image.open(CARPETA_RA / "marcadores" / "codigos" / f"{m['codigo']}.png")
            tmp = carpeta / f"_tmp_{m['id']}.png"
            png.resize((480, 480), Image.NEAREST).save(tmp)
            c.drawImage(str(tmp), cx - lado / 2, cy - lado / 2 + 0.6 * cm, lado, lado)
            tmp.unlink()
        elif img.exists():
            c.drawImage(str(img), cx - lado / 2, cy - lado / 2 + 0.6 * cm, lado, lado, preserveAspectRatio=True)
        c.setDash(3, 3)
        c.setStrokeGray(0.6)
        c.rect(cx - lado / 2 - 0.8 * cm, cy - lado / 2 - 0.6 * cm, lado + 1.6 * cm, lado + 2.0 * cm)
        c.setDash()
        c.setFont("Helvetica-Bold", 26)
        c.setFillGray(0)
        c.drawCentredString(cx, cy - lado / 2 - 0.2 * cm, m["id"])
    c.save()
    return ruta


# ---------------------------------------------------------------- Info del tema
def leer_info(valor: str, avisos: list[str], tema: str) -> list[dict]:
    from convertir import limpiar, texto_inline
    if not valor:
        return []
    if not valor.lower().endswith((".html", ".htm")):
        return [{"t": "p", "texto": limpiar(p)} for p in re.split(r"\n+", valor) if limpiar(p)]
    ruta = CARPETA_RA / "info" / valor
    if not ruta.is_file():
        avisos.append(f"Tema '{tema}': no existe ra/info/{valor}.")
        return []
    sopa = BeautifulSoup(ruta.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for b in sopa.select("script, style"):
        b.decompose()
    bloques = []
    for el in (sopa.body or sopa).find_all(["h2", "h3", "p", "li", "div"]):
        clases = " ".join(el.get("class") or [])
        texto = limpiar(texto_inline(el))
        if not texto:
            continue
        if el.name in ("h2", "h3"):
            bloques.append({"t": "sub", "texto": texto})
        elif el.name == "p" and not el.find_parent("li"):
            bloques.append({"t": "p", "texto": texto})
        elif el.name == "li":
            bloques.append({"t": "item", "texto": texto})
        elif el.name == "div" and re.search(r"equation|ecuacion|formula", clases):
            bloques.append({"t": "ecuacion", "texto": texto})
    return bloques


# ---------------------------------------------------------------- Temas
PATRON_COMB = re.compile(r"^\s*([A-Za-z0-9]+(?:\s*\+\s*[A-Za-z0-9]+)+)\s*=\s*(.+?)\s*$")


def separar_nombre(celda: str) -> tuple[str, str]:
    """'H₂ | h2_m.glb' -> ('H₂', 'h2_m.glb');  'h2_m.glb' -> ('h2_m', 'h2_m.glb')"""
    if "|" in celda:
        nombre, archivo = [p.strip() for p in celda.split("|", 1)]
    else:
        archivo = celda.strip()
        nombre = Path(archivo).stem
    return nombre, archivo


def generar(salida: Path, avisos: list[str]) -> dict | None:
    if not EXCEL_RA.exists():
        return None
    print("\nRealidad Aumentada:")
    wb = load_workbook(EXCEL_RA, data_only=True)
    marcadores, guia = leer_marcadores(wb["Marcadores"], avisos)
    ws = wb["Temas"]

    destino = salida / "ra"
    destino.mkdir(parents=True, exist_ok=True)
    if WEB_RA.is_dir():
        for f in WEB_RA.iterdir():
            if f.is_file():
                shutil.copy2(f, destino / f.name)
    pub = Publicador(destino, avisos)

    encabezados = {_clave(_texto(ws.cell(1, c).value)): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    col_marc = {
        _texto(ws.cell(1, c).value)[len("Marcador "):].strip().upper(): c
        for c in range(1, ws.max_column + 1)
        if _texto(ws.cell(1, c).value).lower().startswith("marcador ")
    }

    def celda(r, nombre):
        c = encabezados.get(nombre)
        return _texto(ws.cell(r, c).value) if c else ""

    temas, categorias, usados, ids = [], [], set(), set()
    for r in range(3, ws.max_row + 1):  # fila 2 = ejemplo
        titulo = celda(r, "tema")
        if not titulo:
            continue
        categoria = celda(r, "categoria") or "Otros"
        if categoria not in categorias:
            categorias.append(categoria)
        tid = _clave(titulo)
        if tid in ids:
            avisos.append(f"Temas, fila {r}: '{titulo}' está repetido — se omite.")
            continue

        modelos = []
        for mid, c in col_marc.items():
            valor = _texto(ws.cell(r, c).value)
            if not valor:
                continue
            if mid not in marcadores:
                avisos.append(f"Temas, fila {r}: el marcador '{mid}' no está en la hoja Marcadores.")
                continue
            nombre, archivo = separar_nombre(valor)
            if not archivo.lower().split("?")[0].endswith((".glb", ".gltf")):
                # Marcador sin modelo (ej. "Medio ácido"): muestra solo su nombre y sirve para combinar
                nombre = valor.split("|")[0].strip()
                modelos.append({"marcador": mid, "nombre": nombre, "archivo": None, "kb": 0, "animado": False})
                usados.add(mid)
                continue
            info = pub.modelo(archivo)
            if info:
                modelos.append({"marcador": mid, "nombre": nombre, **info})
                usados.add(mid)

        combinaciones = []
        for parte in re.split(r"[;\n]+", celda(r, "combinaciones")):
            if not parte.strip():
                continue
            m = PATRON_COMB.match(parte)
            if not m:
                avisos.append(f"Temas, fila {r}: no entiendo la combinación '{parte.strip()}' (usa A+B = archivo.glb).")
                continue
            ids_m = [x.strip().upper() for x in m.group(1).split("+")]
            faltan = [x for x in ids_m if x not in {mo["marcador"] for mo in modelos}]
            if faltan:
                avisos.append(f"Temas, fila {r}: la combinación usa {', '.join(faltan)} pero ese marcador no tiene modelo.")
                continue
            nombre, archivo = separar_nombre(m.group(2))
            info = pub.modelo(archivo)
            if info:
                combinaciones.append({"marcadores": ids_m, "nombre": nombre, **info})

        if not any(m["archivo"] for m in modelos) and not combinaciones:
            avisos.append(f"Temas, fila {r}: '{titulo}' no tiene ningún modelo — se omite.")
            continue

        def numero(nombre, defecto):
            try:
                return float(celda(r, nombre).replace(",", ".")) if celda(r, nombre) else defecto
            except ValueError:
                avisos.append(f"Temas, fila {r}: '{nombre}' no es un número; se usa {defecto}.")
                return defecto

        nobel = celda(r, "nobel-ano")
        ids.add(tid)
        temas.append({
            "id": tid,
            "categoria": categoria,
            "titulo": titulo,
            "descripcion": celda(r, "descripcion") or None,
            "nobel": int(float(nobel)) if nobel.replace(".", "").isdigit() else None,
            "info": leer_info(celda(r, "info"), avisos, titulo),
            "modelos": modelos,
            "combinaciones": combinaciones,
            "distancia": numero("distancia", 2.5),
            "escala": numero("escala", 1.0),
            "girar": _clave(celda(r, "girar") or "si") not in ("no", "0", "false"),
        })

    # Marcadores: imágenes y PDF
    carpeta_m = destino / "marcadores"
    carpeta_m.mkdir(exist_ok=True)
    lista_m = []
    for mid in sorted(marcadores, key=lambda x: (len(x), x)):
        m = dict(marcadores[mid])
        if m["tipo"] == "codigo":
            imagen_marcador_codigo(m["codigo"], mid, carpeta_m / f"{mid}.png")
            m["imagen"] = f"ra/marcadores/{mid}.png"
        else:
            datos, _ = obtener_archivo(m.pop("patronFuente"), "marcadores")
            if datos is None:
                avisos.append(f"No se encontró el patrón del marcador '{mid}'.")
                continue
            (carpeta_m / f"{mid}.patt").write_bytes(datos)
            m["patron"] = f"ra/marcadores/{mid}.patt"
            img_fuente = m.pop("imagenFuente", "")
            if img_fuente:
                datos_img, _ = obtener_archivo(img_fuente, "marcadores")
                if datos_img:
                    (carpeta_m / f"{mid}.png").write_bytes(datos_img)
                    m["imagen"] = f"ra/marcadores/{mid}.png"
        lista_m.append(m)

    if not guia:
        pdf = pdf_marcadores(lista_m, carpeta_m)
        guia = "ra/marcadores/guia_marcadores.pdf" if pdf else ""

    ra = {
        "tipoCodigo": TIPO_CODIGO,
        "guia": guia or None,
        "marcadores": lista_m,
        "categorias": categorias,
        "temas": temas,
    }
    (destino / "temas.json").write_text(json.dumps(ra, ensure_ascii=False), encoding="utf-8")
    total_kb = sum(v["kb"] for v in pub.hechos.values() if v)
    faltan_info = sorted({t["titulo"] for t in temas if not t["info"]})
    if faltan_info:
        print(f"   (sin texto de Info: {len(faltan_info)} temas)")
    print(f"✔ {len(temas)} temas de RA · {len([v for v in pub.hechos.values() if v])} modelos ({total_kb / 1024:.1f} MB en total)")
    return ra
